# -*- coding: utf-8 -*-
"""交接说明里的预算数字（54~100 → 59~115，加入滴速传感器）。可重复执行。"""
import os
from docx import Document

DOC = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   '文档', '项目交接说明.docx')
REPL = [('¥54~100', '¥59~115'), ('54~100', '59~115')]


def main():
    doc = Document(DOC)
    hits = 0
    targets = list(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                targets.extend(cell.paragraphs)
    for p in targets:
        for old, new in REPL:
            if old in p.text:
                for r in p.runs:
                    if old in r.text:
                        r.text = r.text.replace(old, new)
                        hits += 1
    doc.save(DOC)
    print('替换处数:', hits)


if __name__ == '__main__':
    main()
