# -*- coding: utf-8 -*-
"""修正交接说明里的预算数字（¥57~104 → ¥56~103，与采购清单表一致）。"""
import os
from docx import Document

DOC = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   '文档', '项目交接说明.docx')
PAIRS = [('约 ¥57~104', '约 ¥56~103'),
         ('¥57~104', '¥56~103')]


def main():
    doc = Document(DOC)
    hits = 0
    for p in doc.paragraphs:
        for old, new in PAIRS:
            if old in p.text:
                for r in p.runs:
                    if old in r.text:
                        r.text = r.text.replace(old, new)
                        hits += 1
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for old, new in PAIRS:
                        if old in p.text:
                            for r in p.runs:
                                if old in r.text:
                                    r.text = r.text.replace(old, new)
                                    hits += 1
    doc.save(DOC)
    print('replaced runs:', hits)


if __name__ == '__main__':
    main()
