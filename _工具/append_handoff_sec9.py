# -*- coding: utf-8 -*-
"""在《项目交接说明.docx》末尾追加「九、B1~B4 答复与定稿」章节。

原文档是 Letter 纸、微软雅黑 10.5pt，追加内容要跟着它的字体走，
不能套用 502 新文档的宋体样式，否则一份文件里两种字体。
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from docx_common import add_table, add_box

DOC = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   '文档', '项目交接说明.docx')

CN = '微软雅黑'


def retouch(table, size=9.0):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = CN
                    rpr = r._element.get_or_add_rPr()
                    rf = rpr.find(qn('w:rFonts'))
                    if rf is None:
                        rf = OxmlElement('w:rFonts')
                        rpr.insert(0, rf)
                    rf.set(qn('w:ascii'), CN)
                    rf.set(qn('w:hAnsi'), CN)
                    rf.set(qn('w:eastAsia'), CN)
                    if r.font.size is None or r.font.size.pt != size:
                        r.font.size = Pt(size)


def body(doc, text, size=10.5, bold=False, color=None, before=0, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.name = CN
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    rf.set(qn('w:eastAsia'), CN)
    rf.set(qn('w:ascii'), CN)
    rf.set(qn('w:hAnsi'), CN)
    r.font.size = Pt(size)
    r.bold = bold
    if color is not None:
        from docx.shared import RGBColor
        r.font.color.rgb = RGBColor(*color)
    return p


def main():
    doc = Document(DOC)
    n0 = len(doc.tables)

    # 分页
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    doc.add_heading('九、B1~B4 答复与定稿（2026-10-21 晚）', level=1)
    body(doc, '第五节列的 4 个问题（阻塞采购）当晚已答完。本节把答复、定稿结论，以及答复过程中'
              '暴露出来的一个真问题（库存判断错了）一并写清，供后续随时复核。', after=6)

    doc.add_heading('9.1 四项答复', level=2)
    t = add_table(doc, ['#', '问题', '用户答复', '落定结论'],
                  [('B1', 'STM32F103C8T6 还剩几块空闲', 'A. 有 1 块空闲', '主控不买，用闲置那块'),
                   ('B2', 'ESP32-C3 有几个', 'A. 只有 1 个（项目1 占用）', '必须买 1 块同型号'),
                   ('B3', '排母座还剩多少', 'A. 够用', '够；且面包板阶段用不到'),
                   ('B4', '主控走 X1 还是 Z', 'C. 你帮我定', '定为 X1')],
                  [1.0, 6.0, 5.4, 4.6], font_pt=9, aligns=['center', 'left', 'left', 'left'])
    retouch(t)

    doc.add_heading('9.2 选型定稿：X1', level=2)
    body(doc, 'X1 = STM32F103C8T6（复用项目1 框架）+ ESP32-C3（保持 AT 固件）+ STM32 手拼 HTTP 响应。', after=4)
    t = add_table(doc, ['方案', '定稿', '核心理由 / 为什么不是它'],
                  [('X1', '✅ 采用', '复用项目1 已验证的 AT 收发、串口环形缓冲、OLED、按键、蜂鸣器、软件 I2C；'
                                   'Flash 反而更宽（去掉心率血氧/计步约 8KB，新增 HX711+HTTP 约 3KB）'),
                   ('X2', '✗', '要重烧 ESP32 为 Arduino/ESP-IDF，等于两套程序都要自己写，工作量最大'),
                   ('Z', '✗', 'ESP32-C3 单独做主控看起来很省事，但项目1 全部代码作废，且同样要买板子 —— 没有成本优势')],
                  [1.8, 2.0, 13.2], font_pt=9, aligns=['center', 'center', 'left'])
    retouch(t)
    body(doc, '★ X1 的代价照实写：① STM32 里要手工拼 HTTP 报文；② 网页只能是简单静态页（数值 + 自动刷新）。',
         size=9.5, color=(0x60, 0x6B, 0x75), before=2, after=8)

    doc.add_heading('9.3 答复暴露的真问题：库存判断是错的', level=2)
    body(doc, '第四节 4.1 曾把 OLED、蜂鸣器列为"可以直接用"。核对数量后发现在采购清单里它们**各只买了 1 件**，'
              '并且都已装在赖振雄自己的毕设板上 —— 拆下来就毁掉另一个毕设。', after=4)
    t = add_table(doc, ['器件', '项目1 数量', '现状', '本项目结论'],
                  [('ESP32-C3（CORE-ESP32）', '1 块', '已装在自己板上，烧着 AT 固件', '买 1 块新板 + 烧同一份固件'),
                   ('0.96 寸 OLED', '1 个', '已装在自己板上', '买 1 个'),
                   ('有源蜂鸣器模块', '1 套', '已装在自己板上', '买 1 套'),
                   ('红色 LED', '1 个', '已装', '买（分级报警还要黄色）'),
                   ('轻触按键', '6 个（用掉 3）', '剩 3 个', '补 1~2 个（本项目要 4 个）'),
                   ('STM32F103C8T6', '2 块', '1 块闲置', '✅ 用闲置那块')],
                  [4.6, 3.4, 5.2, 3.8], font_pt=9, aligns=['left', 'center', 'left', 'left'])
    retouch(t)
    body(doc, '★ 一个意外的好消息：AT 固件整套文件仍在本地 T:\\TE\\ESP32-C3-MINI-1-AT-V4.1.1.0\\'
              '（esp-at.bin / at_customize.bin / factory_MINI-1.bin / ota_data_initial.bin），'
              '新板烧固件不需要联网。', size=9.5, color=(0x1E, 0x6B, 0x3A), before=2, after=8)

    doc.add_heading('9.4 预算修订', level=2)
    t = add_table(doc, ['项目', '原估算', '修订后', '原因'],
                  [('新购器件合计', '约 ¥15', '约 ¥57~104', 'ESP32/OLED/蜂鸣器/黄灯/按键都要买'),
                   ('新增项：标准砝码', '（原清单没有）', '¥10~25', '没有已知质量基准，±2g 无法被证明'),
                   ('重量传感器量程', '5kg（最常见）', '2kg 全桥', '5kg 的 0.05%FS≈2.5g，单项就吃掉 ±2g 余量')],
                  [3.8, 3.8, 3.8, 5.6], font_pt=9, aligns=['left', 'center', 'center', 'left'])
    retouch(t)

    doc.add_heading('9.5 配套文档（都在 毕业设计2\\文档\\）', level=2)
    t = add_table(doc, ['文档', '作用', '什么时候用'],
                  [('502采购清单(可勾选).xlsx', '12 项采购 + 可复用核对 + 到货验收 + 下单顺序', '现在：下单前'),
                   ('502引脚分配与逐脚连线表.docx', '17 个引脚分配、逐器件连线、12 条硬约束', '到货后：接线时'),
                   ('502可行性验证方案.docx', '8 阶段验收，含称重/温度标定与指标判定表', '接线后：逐阶段验收'),
                   ('项目交接说明.docx（本文）', '课题要求、决策、选型、风险', '随时查阅')],
                  [5.4, 8.2, 3.4], font_pt=9, aligns=['left', 'left', 'left'])
    retouch(t)

    doc.add_heading('9.6 仍需确认（不影响下单）', level=2)
    t = add_table(doc, ['#', '问题', '已按什么假设推进', '若假设不成立怎么办'],
                  [('A4', '输液袋最大多少 ml', '按 ≤1000ml 定 2kg 传感器', '只用 500ml → 换 1kg 更准；3000ml 营养袋 → 必须 5kg'),
                   ('A5', '温度传感器用 DS18B20 还是 PT100', '先用 DS18B20（已有）', '实测超 ±0.5℃ → 换 PT100+MAX31865（约 +¥30）'),
                   ('C1', '交付物范围（只面包板 / 要 PCB）', '先按面包板推进', '要 PCB 再买排母、洞洞板、打板'),
                   ('C2', '时间节点', '未定，按 8 阶段顺序推进', '时间紧则先做 3/4/5 阶段（达标核心），网页最后')],
                  [1.0, 5.2, 5.2, 5.6], font_pt=9, aligns=['center', 'left', 'left', 'left'])
    retouch(t)

    box = add_box(doc, [
        '本方案是「手机连 ESP32 热点」的局域网查看，严格说不等于"上传到平台"。',
        '答辩时若被问到，应主动说明；升级路径是 ESP32 转 STA 连路由器 + AT+MQTTCONN 上巴法云',
        '（项目1 已实测该 AT 固件支持 MQTT），属于二期工作，不在本期交付内。',
    ], fill='FFF6DC', border='BF8F00', title_text='★ 一句必须自己先说清的话')
    retouch(box, size=9)
    for p in doc.paragraphs[-1:]:
        pass

    doc.save(DOC)
    print('SAVED', DOC, '| tables', n0, '->', len(doc.tables), '| paras', len(doc.paragraphs))


if __name__ == '__main__':
    main()
