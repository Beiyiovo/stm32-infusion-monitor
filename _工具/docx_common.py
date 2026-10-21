# -*- coding: utf-8 -*-
"""502 项目文档公共样式库（python-docx）。

约定：正文宋体 10.5pt / 标题黑体 / A4 上下 1.9-1.7cm 左右 2cm。
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BODY_CN = '宋体'
HEAD_CN = '黑体'
MONO = 'Consolas'
NAVY = RGBColor(0x1F, 0x3B, 0x60)
GREY = RGBColor(0x60, 0x6B, 0x75)
RED = RGBColor(0xB0, 0x2A, 0x2A)
GREEN = RGBColor(0x1E, 0x6B, 0x3A)


def _set_style_font(style, size=None, bold=None, color=None, cn=HEAD_CN, latin='Calibri'):
    style.font.name = latin
    rpr = style.element.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr')
        style.element.append(rpr)
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:ascii'), latin)
    rfonts.set(qn('w:hAnsi'), latin)
    rfonts.set(qn('w:eastAsia'), cn)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


def set_run_font(run, size=None, bold=None, cn=BODY_CN, latin='Calibri', color=None, italic=None):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:ascii'), latin)
    rfonts.set(qn('w:hAnsi'), latin)
    rfonts.set(qn('w:eastAsia'), cn)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic


def new_doc():
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)
    rpr = st.element.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr')
        st.element.append(rpr)
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:ascii'), 'Calibri')
    rfonts.set(qn('w:hAnsi'), 'Calibri')
    rfonts.set(qn('w:eastAsia'), BODY_CN)
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.2

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.7)

    for name, size, before, after in (('Heading 1', 15, 14, 6), ('Heading 2', 12, 10, 4), ('Heading 3', 11, 8, 3)):
        _set_style_font(doc.styles[name], size=size, bold=True, color=NAVY)
        p = doc.styles[name].paragraph_format
        p.space_before, p.space_after = Pt(before), Pt(after)
        p.line_spacing = 1.15
        p.keep_with_next = True
    return doc


def title(doc, main, sub=None, meta=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(main)
    set_run_font(r, size=20, bold=True, cn=HEAD_CN, color=NAVY)
    if sub:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(sub)
        set_run_font(r2, size=11.5, bold=False, cn=BODY_CN, color=GREY)
    if meta:
        p3 = doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(10)
        r3 = p3.add_run(meta)
        set_run_font(r3, size=9, cn=BODY_CN, color=GREY)
        # 底部分隔线
        pPr = p3._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '4')
        bottom.set(qn('w:color'), 'BFBFBF')
        pBdr.append(bottom)
        pPr.append(pBdr)


def h1(doc, text):
    return doc.add_heading(text, level=1)


def h2(doc, text):
    return doc.add_heading(text, level=2)


def h3(doc, text):
    return doc.add_heading(text, level=3)


def para(doc, text, size=10.5, bold=False, color=None, align=None, space_after=4, space_before=0, indent=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if indent is not None:
        pf.first_line_indent = Cm(indent)
    if align == 'center':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'right':
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    return p


def rich(doc, chunks, size=10.5, space_after=4, align=None):
    """chunks: [(text, {bold/color/size/cn}), ...]"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align == 'center':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text, opt in chunks:
        r = p.add_run(text)
        set_run_font(r,
                     size=opt.get('size', size),
                     bold=opt.get('bold', False),
                     color=opt.get('color'),
                     cn=opt.get('cn', BODY_CN),
                     latin=opt.get('latin', 'Calibri'))
    return p


def bullets(doc, items, size=10.5, style='List Bullet', space_after=2):
    for it in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(it)
        set_run_font(r, size=size)


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def _cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for edge, w in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:w'), str(w))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tblPr.append(mar)


def _borders(table, color='9AA7B4', outer='8', inner='4'):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), outer if edge in ('top', 'bottom', 'left', 'right') else inner)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        b.append(el)
    tblPr.append(b)


def _no_split(row):
    """禁止单行跨页被拆开（否则会在下一页留一个空行）。"""
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


def _set_widths(table, widths_cm):
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    for row in table.rows:
        for idx, w in enumerate(widths_cm):
            row.cells[idx].width = Cm(w)
    # 列宽也要写到 tblGrid，否则 WPS/Word 会重新分配
    grid = table._tbl.find(qn('w:tblGrid'))
    if grid is not None:
        for idx, gc in enumerate(grid.findall(qn('w:gridCol'))):
            if idx < len(widths_cm):
                gc.set(qn('w:w'), str(int(widths_cm[idx] * 567)))


def add_table(doc, headers, rows, widths, font_pt=8.5, header_fill='1F3B60',
              zebra='F2F5F9', first_col_bold=False, aligns=None, row_fills=None):
    n = len(headers)
    table = doc.add_table(rows=1, cols=n)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(table)
    _cell_margins(table)

    hdr = table.rows[0]
    _repeat_header(hdr)
    _no_split(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _shade(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=font_pt, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    for ridx, row in enumerate(rows):
        trow = table.add_row()
        _no_split(trow)
        cells = trow.cells
        fill = None
        if row_fills and ridx < len(row_fills) and row_fills[ridx]:
            fill = row_fills[ridx]
        elif zebra and ridx % 2 == 1:
            fill = zebra
        for i, text in enumerate(row):
            cell = cells[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if fill:
                _shade(cell, fill)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            a = (aligns[i] if aligns and i < len(aligns) else 'left')
            p.alignment = {'center': WD_ALIGN_PARAGRAPH.CENTER,
                           'right': WD_ALIGN_PARAGRAPH.RIGHT}.get(a, WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(str(text))
            set_run_font(r, size=font_pt, bold=(first_col_bold and i == 0))

    _set_widths(table, widths)
    return table


def add_box(doc, lines, fill='FFF6DC', border='BF8F00', title_text=None, font_pt=9.5, width_cm=17.0):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '8')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), border)
        b.append(el)
    tblPr.append(b)
    _cell_margins(table, top=90, bottom=90, left=140, right=140)
    cell = table.rows[0].cells[0]
    _shade(cell, fill)
    p0 = cell.paragraphs[0]
    first = True
    if title_text:
        p0.paragraph_format.space_after = Pt(2)
        r = p0.add_run(title_text)
        set_run_font(r, size=font_pt, bold=True, color=RGBColor(0x7F, 0x5F, 0x00), cn=HEAD_CN)
        first = False
    for line in lines:
        p = p0 if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(line)
        set_run_font(r, size=font_pt)
    _set_widths(table, [width_cm])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_code(doc, lines, width_cm=17.0, font_pt=8.5):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(table, color='C9D1D9')
    _cell_margins(table, top=80, bottom=80, left=140, right=100)
    cell = table.rows[0].cells[0]
    _shade(cell, 'F5F6F8')
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, size=font_pt, cn=MONO, latin=MONO)
    _set_widths(table, [width_cm])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def _field(p, instr):
    r1 = p.add_run()
    fc = OxmlElement('w:fldChar')
    fc.set(qn('w:fldCharType'), 'begin')
    r1._r.append(fc)
    r2 = p.add_run()
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = instr
    r2._r.append(it)
    r3 = p.add_run()
    fc2 = OxmlElement('w:fldChar')
    fc2.set(qn('w:fldCharType'), 'end')
    r3._r.append(fc2)
    for r in (r1, r2, r3):
        set_run_font(r, size=8, color=GREY)


def add_footer(doc, text):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text + '   ·   第 ')
    set_run_font(r, size=8, color=GREY)
    _field(p, 'PAGE')
    r2 = p.add_run(' 页')
    set_run_font(r2, size=8, color=GREY)


def page_break(doc):
    from docx.enum.text import WD_BREAK
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
